#!/usr/bin/env python3
# coding: utf-8
# SPDX-License-Identifier: MIT
"""Extract one revised HMD sheet without modifying its source workbook."""

from __future__ import annotations

import argparse
import csv
import hashlib
import os
import sys
import tempfile
from pathlib import Path
from typing import Sequence

from openpyxl import load_workbook


HEADER = [
    "sequence",
    "module",
    "level",
    "type",
    "identifier",
    "name",
    "datatype",
    "multiplicity",
    "domain_name",
    "definition",
    "label_local",
    "definition_local",
    "element",
    "id",
    "semantic_path",
    "class_term",
]


class HmdWorkbookError(ValueError):
    """The requested sheet is not a safe revised-HMD input."""


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def text(value: object) -> str:
    if value is None:
        return ""
    return str(value)


def extract(workbook_path: Path, sheet_name: str) -> tuple[list[list[str]], str]:
    if not workbook_path.is_file():
        raise HmdWorkbookError(f"workbook does not exist: {workbook_path}")
    before = sha256(workbook_path)
    workbook = load_workbook(
        workbook_path,
        read_only=False,
        data_only=False,
        keep_links=True,
    )
    try:
        matches = [sheet for sheet in workbook.worksheets if sheet.title == sheet_name]
        if len(matches) != 1:
            raise HmdWorkbookError(
                f"expected exactly one sheet {sheet_name!r}, got {len(matches)}"
            )
        sheet = matches[0]
        if workbook._external_links:
            raise HmdWorkbookError("external workbook links are not allowed")
        if sheet.merged_cells.ranges:
            raise HmdWorkbookError(f"{sheet_name}: merged cells are not allowed")
        hidden_rows = [index for index, item in sheet.row_dimensions.items() if item.hidden]
        hidden_columns = [
            name for name, item in sheet.column_dimensions.items() if item.hidden
        ]
        if hidden_rows or hidden_columns:
            raise HmdWorkbookError(
                f"{sheet_name}: hidden rows/columns are not allowed: "
                f"rows={hidden_rows!r}, columns={hidden_columns!r}"
            )

        actual = [text(sheet.cell(1, column).value) for column in range(1, 17)]
        if actual != HEADER or sheet.max_column != len(HEADER):
            raise HmdWorkbookError(
                f"{sheet_name}: HMD header mismatch; expected {HEADER!r}, "
                f"got {actual!r} with {sheet.max_column} column(s)"
            )

        rows: list[list[str]] = []
        for row_number in range(2, sheet.max_row + 1):
            values = [
                sheet.cell(row_number, column).value
                for column in range(1, len(HEADER) + 1)
            ]
            if not any(value not in (None, "") for value in values):
                continue
            formulas = [
                HEADER[index]
                for index, value in enumerate(values)
                if isinstance(value, str) and value.startswith("=")
            ]
            if formulas:
                raise HmdWorkbookError(
                    f"{sheet_name}:{row_number}: formulas are not allowed in "
                    f"{formulas!r}"
                )
            rows.append([text(value) for value in values])
    finally:
        workbook.close()
    after = sha256(workbook_path)
    if before != after:
        raise HmdWorkbookError(
            f"source workbook changed while reading: before={before}, after={after}"
        )
    if not rows:
        raise HmdWorkbookError(f"{sheet_name}: HMD has no data rows")
    return rows, before


def write_csv(output_path: Path, rows: list[list[str]], encoding: str) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary_name = tempfile.mkstemp(
        prefix=f".{output_path.name}.",
        suffix=".tmp",
        dir=output_path.parent,
    )
    os.close(descriptor)
    temporary = Path(temporary_name)
    try:
        with temporary.open("w", encoding=encoding, newline="") as handle:
            writer = csv.writer(handle, lineterminator="\n")
            writer.writerow(HEADER)
            writer.writerows(rows)
        os.replace(temporary, output_path)
    finally:
        if temporary.exists():
            temporary.unlink()


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Extract an exact revised 16-column HMD sheet to CSV"
    )
    parser.add_argument("workbook", type=Path)
    parser.add_argument("sheet")
    parser.add_argument("output_csv", type=Path)
    parser.add_argument("-e", "--encoding", default="utf-8-sig")
    return parser


def main(argv: Sequence[str] | None = None) -> int:
    args = build_parser().parse_args(argv)
    try:
        rows, workbook_sha256 = extract(args.workbook, args.sheet)
        write_csv(args.output_csv, rows, args.encoding)
    except (OSError, HmdWorkbookError) as exc:
        print(f"extract_hmd_workbook.py: error: {exc}", file=sys.stderr)
        return 2
    print(
        f"Wrote {len(rows)} HMD row(s) from {args.sheet!r} to "
        f"{args.output_csv}; workbook_sha256={workbook_sha256}"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
